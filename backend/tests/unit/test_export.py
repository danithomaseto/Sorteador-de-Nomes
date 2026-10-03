import codecs
import csv
import io
from datetime import UTC, datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.draw_engine import WITH_REPETITION, WITHOUT_REPETITION
from app.exporting import ExportDocument, ExportRound, ExportWinner, write_csv, write_xlsx
from app.exporting.document import describe_algorithm, slugify
from app.exporting.safe_cells import neutralize_formula

SAO_PAULO = ZoneInfo("America/Sao_Paulo")


def make_round(
    number: int = 1, winners: tuple[str, ...] = ("Maria Silva", "João Souza")
) -> ExportRound:
    return ExportRound(
        number=number,
        drawn_at=datetime(2026, 10, 3, 21, 35, 12, tzinfo=UTC),
        quantity=len(winners),
        allow_repeat=False,
        remove_winners=True,
        total_participants=100,
        pool_size=95,
        algorithm=f"{WITHOUT_REPETITION}+os-csprng",
        winners=tuple(ExportWinner(position=i, name=n) for i, n in enumerate(winners, start=1)),
    )


def make_document(*rounds: ExportRound, name: str = "Churrasco da equipe") -> ExportDocument:
    return ExportDocument(
        draw_name=name,
        timezone=SAO_PAULO,
        generated_at=datetime(2026, 10, 3, 21, 40, tzinfo=UTC),
        rounds=rounds or (make_round(),),
    )


def read_csv(data: bytes) -> list[list[str]]:
    assert data.startswith(codecs.BOM_UTF8)
    text = data[len(codecs.BOM_UTF8) :].decode("utf-8")
    return list(csv.reader(io.StringIO(text), delimiter=";"))


def sheet_values(data: bytes) -> list[tuple[object, ...]]:
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


class TestCsv:
    def test_cabecalho_e_linhas(self) -> None:
        rows = read_csv(write_csv(make_document()))
        assert rows[0][:6] == [
            "Sorteio",
            "Rodada",
            "Data e hora",
            "Fuso horário",
            "Posição",
            "Vencedor",
        ]
        assert rows[1] == [
            "Churrasco da equipe",
            "1",
            "03/10/2026 18:35:12",
            "America/Sao_Paulo",
            "1",
            "Maria Silva",
            "Não",
            "Sim",
            "95",
            "100",
        ]
        assert rows[2][5] == "João Souza"

    def test_varias_rodadas(self) -> None:
        document = make_document(make_round(1, ("Ana",)), make_round(2, ("Bia", "Caio")))
        rows = read_csv(write_csv(document))
        assert [(row[1], row[5]) for row in rows[1:]] == [("1", "Ana"), ("2", "Bia"), ("2", "Caio")]

    def test_neutraliza_formulas(self) -> None:
        document = make_document(
            make_round(winners=("=HYPERLINK(1)", "@SOMA", "+1", "-2", "Ana")), name="=cmd"
        )
        rows = read_csv(write_csv(document))
        assert [row[5] for row in rows[1:]] == ["'=HYPERLINK(1)", "'@SOMA", "'+1", "'-2", "Ana"]
        assert rows[1][0] == "'=cmd"


class TestXlsx:
    def test_uma_rodada_com_metadados(self) -> None:
        values = sheet_values(write_xlsx(make_document()))
        flat = {row[0]: row[1] for row in values if row and row[0] is not None}
        assert values[0][0] == "Resultado do sorteio"
        assert flat["Sorteio"] == "Churrasco da equipe"
        assert flat["Data e hora da rodada"] == "03/10/2026 18:35:12"
        assert flat["Vencedores sorteados"] == 2
        assert flat["Participantes disponíveis na rodada"] == 95
        assert flat["Repetição na mesma rodada"] == "Não"
        assert flat["Vencedores removidos das próximas rodadas"] == "Sim"
        assert "Fisher–Yates" in str(flat["Método"])
        assert (1, "Maria Silva") in values
        assert (2, "João Souza") in values

    def test_varias_rodadas_em_tabela(self) -> None:
        document = make_document(make_round(1, ("Ana",)), make_round(2, ("Bia",)))
        values = sheet_values(write_xlsx(document))
        header = next(row for row in values if row[0] == "Rodada")
        assert header[:4] == ("Rodada", "Data e hora", "Posição", "Vencedor")
        assert tuple(values[values.index(header) + 2][:4]) == (2, "03/10/2026 18:35:12", 1, "Bia")

    def test_formula_vira_texto(self) -> None:
        data = write_xlsx(make_document(make_round(winners=("=1+1",)), name="=cmd"))
        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        cells = [
            cell for row in sheet.iter_rows() for cell in row if cell.value in {"=1+1", "=cmd"}
        ]
        assert len(cells) == 2
        assert all(cell.data_type == "s" for cell in cells)


class TestTextos:
    def test_nome_do_arquivo(self) -> None:
        assert make_document().file_stem() == "churrasco-da-equipe-rodada-1"
        two = make_document(make_round(1), make_round(2), name="Sorteio Ação!")
        assert two.file_stem() == "sorteio-acao-rodadas"
        assert make_document(name="???").file_stem() == "sorteio-rodada-1"

    def test_slug_limita_tamanho(self) -> None:
        assert len(slugify("a" * 300)) == 40

    def test_descricao_do_algoritmo(self) -> None:
        assert "criptográfico" in describe_algorithm(f"{WITHOUT_REPETITION}+os-csprng")
        assert "com reposição" in describe_algorithm(f"{WITH_REPETITION}+os-csprng")

    def test_neutraliza_apenas_gatilhos(self) -> None:
        assert neutralize_formula("Ana") == "Ana"
        assert neutralize_formula("\tAna") == "'\tAna"
