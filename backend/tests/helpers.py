"""Geradores de arquivos para os testes (nenhum arquivo real de pessoas é usado)."""

import io
import zipfile
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.importing import ImportLimits

MEGABYTE = 1024 * 1024

DEFAULT_LIMITS = ImportLimits(
    max_bytes=5 * MEGABYTE,
    max_uncompressed_bytes=50 * MEGABYTE,
    max_rows=50_000,
    max_text_chars=2_000_000,
)

type Rows = Sequence[Sequence[object]]


def make_xlsx(
    rows: Rows | None = None,
    *,
    sheets: dict[str, Rows] | None = None,
    hidden: frozenset[str] = frozenset(),
    merge: str | None = None,
) -> bytes:
    """Cria um .xlsx em memória. ``rows`` gera uma aba única chamada "Plan1"."""
    workbook = Workbook()
    default = workbook.active
    assert isinstance(default, Worksheet)
    workbook.remove(default)
    for title, sheet_rows in (sheets or {"Plan1": rows or []}).items():
        sheet = workbook.create_sheet(title)
        for row in sheet_rows:
            sheet.append(list(row))
        if title in hidden:
            sheet.sheet_state = "hidden"
        if merge:
            sheet.merge_cells(merge)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_zip(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def large_xlsx(count: int) -> bytes:
    """Planilha "Nome | Matrícula | Área" com ``count`` participantes (modo de escrita rápida)."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Participantes")
    sheet.append(["Nome", "Matrícula", "Área"])
    for index in range(count):
        sheet.append([f"Participante {index:05d}", 10_000 + index, "Operações"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
