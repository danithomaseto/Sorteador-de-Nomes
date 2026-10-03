import io
from typing import Any, get_args

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.draw_engine import WITH_REPETITION, WITHOUT_REPETITION
from app.schemas.exports import Algorithm
from tests.api.conftest import ClientFactory


def export_body(**overrides: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "draw_name": "Churrasco da equipe",
        "timezone": "America/Sao_Paulo",
        "rounds": [
            {
                "number": 1,
                "drawn_at": "2026-10-03T21:35:12Z",
                "quantity": 2,
                "allow_repeat": False,
                "remove_winners": True,
                "total_participants": 127,
                "pool_size": 127,
                "algorithm": f"{WITHOUT_REPETITION}+os-csprng",
                "winners": [
                    {"position": 1, "name": "Maria Silva"},
                    {"position": 2, "name": "João"},
                ],
            }
        ],
    }
    body.update(overrides)
    return body


def test_algoritmos_aceitos_acompanham_o_motor() -> None:
    assert set(get_args(Algorithm.__value__)) == {
        f"{WITHOUT_REPETITION}+os-csprng",
        f"{WITH_REPETITION}+os-csprng",
    }


def test_xlsx(client: TestClient) -> None:
    response = client.post("/api/v1/exports?format=xlsx", json=export_body())
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="churrasco-da-equipe-rodada-1.xlsx"'
    )
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    values = list(sheet.iter_rows(values_only=True))
    assert (1, "Maria Silva") in values
    assert ("Data e hora da rodada", "03/10/2026 18:35:12") in [row[:2] for row in values]


def test_csv(client: TestClient) -> None:
    response = client.post("/api/v1/exports?format=csv", json=export_body())
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    text = response.content.decode("utf-8-sig")
    assert "Churrasco da equipe;1;03/10/2026 18:35:12;America/Sao_Paulo;1;Maria Silva" in text


def test_nome_do_sorteio_vazio_vira_padrao(client: TestClient) -> None:
    response = client.post("/api/v1/exports?format=csv", json=export_body(draw_name="   "))
    assert response.content.decode("utf-8-sig").splitlines()[1].startswith("Sorteio;1;")


def test_fuso_invalido(client: TestClient) -> None:
    for timezone in ["Marte/Base", "../../etc/passwd", ""]:
        response = client.post("/api/v1/exports", json=export_body(timezone=timezone))
        assert response.status_code == 422, timezone
        assert response.json()["code"] == "invalid_timezone"


def test_algoritmo_desconhecido(client: TestClient) -> None:
    body = export_body()
    body["rounds"][0]["algorithm"] = "manual"
    response = client.post("/api/v1/exports", json=body)
    assert response.status_code == 422
    assert response.json()["errors"][0]["field"] == "rounds.0.algorithm"


def test_nome_de_vencedor_invalido(client: TestClient) -> None:
    body = export_body()
    body["rounds"][0]["winners"][0]["name"] = "x" * 121
    response = client.post("/api/v1/exports", json=body)
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_name"


def test_data_sem_fuso_e_recusada(client: TestClient) -> None:
    body = export_body()
    body["rounds"][0]["drawn_at"] = "2026-10-03T21:35:12"
    response = client.post("/api/v1/exports", json=body)
    assert response.status_code == 422


def test_rodadas_demais(make_client: ClientFactory) -> None:
    client = make_client(max_export_rounds=1)
    body = export_body()
    body["rounds"] = body["rounds"] * 2
    response = client.post("/api/v1/exports", json=body)
    assert response.json()["params"] == {"limit": "export_rounds", "maximum": 1}
