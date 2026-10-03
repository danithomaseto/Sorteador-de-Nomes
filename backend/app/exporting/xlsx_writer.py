"""Planilha .xlsx legível: cabeçalho com os metadados do sorteio e tabela de vencedores."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.exporting.document import ExportDocument, ExportRound, describe_algorithm, yes_no

_TITLE = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="F2B300")
_NOTE = Font(italic=True, color="5F5D57")
_NOTE_TEXT = (
    "Arquivo gerado pelo Sorteia. Os dados do sorteio não ficam armazenados no serviço: "
    "guarde este arquivo se precisar de um registro."
)


def write_xlsx(document: ExportDocument) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):  # pragma: no cover - um Workbook novo sempre tem uma aba
        raise TypeError("planilha sem aba ativa")
    sheet.title = "Resultado"

    title = _cell(sheet, 1, 1)
    title.value = "Resultado do sorteio"
    title.font = _TITLE
    row = 3
    single = document.rounds[0] if len(document.rounds) == 1 else None
    metadata: list[tuple[str, str | int]] = [
        ("Sorteio", document.draw_name),
        ("Rodadas", ", ".join(str(r.number) for r in document.rounds)),
        ("Fuso horário", document.timezone.key),
        ("Arquivo gerado em", document.local(document.generated_at)),
    ]
    if single is not None:
        metadata += _round_metadata(document, single)
    for label, value in metadata:
        _label(_cell(sheet, row, 1), label)
        _value(_cell(sheet, row, 2), value)
        row += 1

    row += 1
    header_row = row
    if single is not None:
        _table_header(sheet, row, ("Posição", "Vencedor"))
        for winner in single.winners:
            row += 1
            _cell(sheet, row, 1).value = winner.position
            _value(_cell(sheet, row, 2), winner.name)
    else:
        _table_header(
            sheet,
            row,
            (
                "Rodada",
                "Data e hora",
                "Posição",
                "Vencedor",
                "Repetição",
                "Removidos",
                "Disponíveis",
            ),
        )
        for draw_round in document.rounds:
            for winner in draw_round.winners:
                row += 1
                values: tuple[str | int, ...] = (
                    draw_round.number,
                    document.local(draw_round.drawn_at),
                    winner.position,
                    winner.name,
                    yes_no(draw_round.allow_repeat),
                    yes_no(draw_round.remove_winners),
                    draw_round.pool_size,
                )
                for column, value in enumerate(values, start=1):
                    _value(_cell(sheet, row, column), value)

    note = _cell(sheet, row + 2, 1)
    note.value = _NOTE_TEXT
    note.font = _NOTE

    sheet.freeze_panes = f"A{header_row + 1}"
    for letter, width in zip("ABCDEFG", (28, 48, 12, 40, 12, 12, 14), strict=True):
        sheet.column_dimensions[letter].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _round_metadata(
    document: ExportDocument, draw_round: ExportRound
) -> list[tuple[str, str | int]]:
    return [
        ("Data e hora da rodada", document.local(draw_round.drawn_at)),
        ("Vencedores sorteados", draw_round.quantity),
        ("Participantes disponíveis na rodada", draw_round.pool_size),
        ("Total de participantes na lista", draw_round.total_participants),
        ("Repetição na mesma rodada", yes_no(draw_round.allow_repeat)),
        ("Vencedores removidos das próximas rodadas", yes_no(draw_round.remove_winners)),
        ("Método", describe_algorithm(draw_round.algorithm)),
    ]


def _cell(sheet: Worksheet, row: int, column: int) -> Cell:
    cell = sheet.cell(row=row, column=column)
    if not isinstance(cell, Cell):  # pragma: no cover - não mesclamos células
        raise TypeError("célula mesclada inesperada")
    return cell


def _label(cell: Cell, text: str) -> None:
    cell.value = text
    cell.font = _BOLD


def _value(cell: Cell, value: str | int) -> None:
    cell.value = value
    if isinstance(value, str):
        # Tipo "string" explícito: um texto que começa com "=" nunca vira fórmula.
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top")


def _table_header(sheet: Worksheet, row: int, titles: tuple[str, ...]) -> None:
    for column, title in enumerate(titles, start=1):
        cell = _cell(sheet, row, column)
        _label(cell, title)
        cell.fill = _HEADER_FILL
