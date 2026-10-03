"""Leitura segura de planilhas .xlsx.

Proteções, nesta ordem:

1. assinatura do arquivo: ZIP é aceito; CFB (``.xls`` antigo ou planilha protegida por senha)
   recebe mensagem própria; qualquer outra coisa é recusada;
2. estrutura OOXML mínima (``[Content_Types].xml`` e ``xl/workbook.xml``) — recusa .docx, .ods etc.;
3. tamanho total descompactado declarado no ZIP (o ``zipfile`` nunca entrega mais bytes que o
   declarado, então o limite vale também contra zip bombs);
4. ``openpyxl`` em modo somente leitura (streaming) e com ``defusedxml`` instalado, que bloqueia
   ataques de expansão de entidades XML; fórmulas viram o último valor calculado;
5. leitura interrompida ao passar do limite de linhas ou após muitas linhas vazias seguidas.

Tudo acontece em memória; nada é gravado em disco.
"""

import logging
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Protocol, cast

import openpyxl

from app.importing.cells import CellValue
from app.importing.errors import (
    ImportFailedError,
    InvalidSpreadsheetError,
    LegacyOrProtectedSpreadsheetError,
    SheetNotFoundError,
    SpreadsheetTooLargeError,
    TooManyRowsError,
    UnsupportedFileTypeError,
)
from app.importing.limits import ImportLimits

logger = logging.getLogger(__name__)

ZIP_SIGNATURE = b"PK\x03\x04"
CFB_SIGNATURE = bytes.fromhex("d0cf11e0a1b11ae1")
_REQUIRED_PARTS = frozenset({"[Content_Types].xml", "xl/workbook.xml"})
_MAX_ZIP_ENTRIES = 10_000
_DATA_PROBE_ROWS = 200


class _ReadOnlySheet(Protocol):
    """O que usamos de ``openpyxl`` em modo somente leitura (os stubs não modelam esse modo)."""

    title: str
    sheet_state: str

    def reset_dimensions(self) -> None: ...

    def iter_rows(
        self, *, max_row: int | None = None, values_only: bool = True
    ) -> Iterator[tuple[object, ...]]: ...


@dataclass(frozen=True, slots=True)
class SheetInfo:
    index: int
    name: str
    hidden: bool


@dataclass(frozen=True, slots=True)
class WorkbookData:
    sheets: tuple[SheetInfo, ...]
    selected: int
    rows: list[tuple[CellValue, ...]]


def read_workbook(data: bytes, limits: ImportLimits, sheet: int | None) -> WorkbookData:
    _check_container(data, limits)
    try:
        workbook = openpyxl.load_workbook(
            BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
    except Exception as exc:
        # Fronteira com um parser de arquivo não confiável: qualquer falha vira erro amigável.
        logger.info("xlsx_unreadable", extra={"exc_type": type(exc).__name__})
        raise InvalidSpreadsheetError from exc
    try:
        worksheets = [cast(_ReadOnlySheet, ws) for ws in workbook.worksheets]
        if not worksheets:
            raise InvalidSpreadsheetError
        sheets = tuple(
            SheetInfo(index=i, name=ws.title, hidden=ws.sheet_state != "visible")
            for i, ws in enumerate(worksheets)
        )
        selected = _first_sheet_with_data(worksheets, sheets) if sheet is None else sheet
        if not 0 <= selected < len(worksheets):
            raise SheetNotFoundError(selected)
        rows = _read_rows(worksheets[selected], limits)
    except ImportFailedError:
        raise
    except Exception as exc:
        logger.info("xlsx_unreadable", extra={"exc_type": type(exc).__name__})
        raise InvalidSpreadsheetError from exc
    finally:
        workbook.close()
    return WorkbookData(sheets=sheets, selected=selected, rows=rows)


def _check_container(data: bytes, limits: ImportLimits) -> None:
    if data.startswith(CFB_SIGNATURE):
        raise LegacyOrProtectedSpreadsheetError
    if not data.startswith(ZIP_SIGNATURE):
        raise InvalidSpreadsheetError
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise InvalidSpreadsheetError from exc
    if len(entries) > _MAX_ZIP_ENTRIES:
        raise SpreadsheetTooLargeError
    if not {entry.filename for entry in entries} >= _REQUIRED_PARTS:
        raise UnsupportedFileTypeError
    if sum(entry.file_size for entry in entries) > limits.max_uncompressed_bytes:
        raise SpreadsheetTooLargeError


def _first_sheet_with_data(
    worksheets: Sequence[_ReadOnlySheet], sheets: tuple[SheetInfo, ...]
) -> int:
    """Primeira aba (visíveis antes das ocultas) com algum conteúdo nas primeiras linhas."""
    for info in sorted(sheets, key=lambda item: item.hidden):
        worksheet = worksheets[info.index]
        worksheet.reset_dimensions()
        for row in worksheet.iter_rows(max_row=_DATA_PROBE_ROWS, values_only=True):
            if any(_has_content(value) for value in row):
                return info.index
    return 0


def _read_rows(worksheet: _ReadOnlySheet, limits: ImportLimits) -> list[tuple[CellValue, ...]]:
    # As dimensões gravadas no arquivo podem estar erradas (alguns geradores gravam "A1:A1");
    # sem elas, o openpyxl lê todas as células realmente presentes.
    worksheet.reset_dimensions()
    rows: list[tuple[CellValue, ...]] = []
    non_empty = 0
    blank_streak = 0
    for raw_row in worksheet.iter_rows(values_only=True):
        row = tuple(_to_cell_value(value) for value in raw_row[: limits.max_columns])
        if not any(_has_content(value) for value in row):
            blank_streak += 1
            if blank_streak > limits.max_blank_streak:
                break
            rows.append(())
            continue
        blank_streak = 0
        non_empty += 1
        if non_empty > limits.max_rows + 1:  # +1: a primeira linha pode ser cabeçalho
            raise TooManyRowsError(limits.max_rows)
        rows.append(row)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def _to_cell_value(value: object) -> CellValue:
    """Converte o que o openpyxl devolve nos tipos que a importação entende."""
    if value is None or isinstance(value, str | int | float | date | time | timedelta):
        return value  # bool é int; datetime é date
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def _has_content(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True
